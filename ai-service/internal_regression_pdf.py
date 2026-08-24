"""
internal_regression_pdf.py
Rapport PDF + XLSX pour les tests de Regression Interne (Internal Regression).
Réutilise les mêmes execution_results que le regression classique
(regression_generator.py / regression_runner.py) — name, category, severity,
page, url, status, reason, duration, priority, expected, description — juste
habillé "Internal" avec les infos d'auth (cookies/token/credentials), comme
internal_api_pdf.py le fait pour l'API.
"""

from ctypes.macholib import framework
from io import BytesIO
from datetime import datetime
import math

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.units import mm

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference


# ── Réutilise les helpers déjà définis dans pdf_generator.py ────────────────
from pdf_generator import (
    NAVY, GOLD, RED, RED_BG, GREEN, GREEN_BG, ORANGE, ORANGE_BG,
    BLUE, BLUE_BG, LIGHT_BG, BORDER, BORDER_DARK, MUTED, WHITE, INDIGO, INDIGO_BG,
    TEAL, PURPLE, PRIORITY_COLORS,
    section_header, sub_section_header, stat_card,
    _grade_from_score, _make_score_gauge, _build_certificate_card,
)

ACCENT = HexColor('#eab308')  # jaune

CAT_COLORS = {
    'authentication': '#6366f1',
    'navigation':      '#10b981',
    'content':         '#3b82f6',
    'functionality':   '#8b5cf6',
}
CAT_ORDER = ['Authentication', 'Navigation', 'Content', 'Functionality']

SEV_COLORS = {
    'critical': '#ef4444',
    'high':     '#f97316',
    'medium':   '#f59e0b',
    'low':      '#10b981',
}


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _extract_inputs(generation_data: dict):
    tests = (
        generation_data.get('execution_results') or
        generation_data.get('test_cases') or
        generation_data.get('result', {}).get('execution_results') or
        generation_data.get('result', {}).get('test_cases') or
        []
    )
    ai_data = (
        generation_data.get('ai') or
        generation_data.get('result', {}).get('ai') or
        generation_data.get('_groq_recs') or
        {}
    )
    return tests, ai_data


def _base_url(generation_data: dict) -> str:
    return generation_data.get('url') or generation_data.get('result', {}).get('url', '')


def _auth_method(generation_data: dict) -> str:
    if generation_data.get('token') or generation_data.get('result', {}).get('token'):
        return 'Token (JWT)'
    if generation_data.get('cookies') or generation_data.get('result', {}).get('cookies'):
        return 'Cookies'
    if generation_data.get('username') or generation_data.get('result', {}).get('username'):
        return 'Credentials (username/password)'
    return 'N/A'


def _category_label(t: dict) -> str:
    cat = (t.get('category') or 'navigation').lower()
    return {
        'authentication': 'Authentication', 'navigation': 'Navigation',
        'content': 'Content', 'functionality': 'Functionality',
    }.get(cat, cat.title() if cat else 'Navigation')


def _compute_score(tests: list) -> tuple:
    """Score pondéré — authentication + severity critical/high comptent plus."""
    if not tests:
        return 0, 'Critical', '#ef4444'
    HIGH_CATS = {'authentication'}
    total_w = 0
    earned_w = 0
    for t in tests:
        cat = (t.get('category') or 'navigation').lower()
        sev = (t.get('severity') or t.get('priority') or 'medium').lower()
        w = 3 if (cat in HIGH_CATS or sev in ('critical', 'high')) else 1
        total_w += w
        if t.get('status') == 'pass':
            earned_w += w
    score = round(earned_w / total_w * 100) if total_w else 0
    if score >= 90:
        return score, 'Excellent', '#10b981'
    if score >= 75:
        return score, 'Good', '#22c55e'
    if score >= 50:
        return score, 'Acceptable', '#f59e0b'
    return score, 'Critical', '#ef4444'


def _is_critical_fail(t: dict) -> bool:
    if t.get('status') != 'fail':
        return False
    cat = (t.get('category') or '').lower()
    sev = (t.get('severity') or t.get('priority') or '').lower()
    return cat == 'authentication' or sev in ('critical', 'high')


# ─────────────────────────────────────────────────────────────────────────────
# Chart (matplotlib) — category breakdown
# ─────────────────────────────────────────────────────────────────────────────

def _make_category_chart(tests: list):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    from io import BytesIO as BIO
    from reportlab.platypus import Image as RLImage

    cats = {c: {'pass': 0, 'fail': 0} for c in CAT_ORDER}
    for t in tests:
        label = _category_label(t)
        if label not in cats:
            cats[label] = {'pass': 0, 'fail': 0}
        if t.get('status') == 'pass':
            cats[label]['pass'] += 1
        elif t.get('status') == 'fail':
            cats[label]['fail'] += 1

    labels = [c for c in CAT_ORDER if cats.get(c, {}).get('pass', 0) + cats.get(c, {}).get('fail', 0) > 0]
    if not labels:
        labels = CAT_ORDER
    p = [cats[c]['pass'] for c in labels]
    f = [cats[c]['fail'] for c in labels]

    plt.rcParams.update({'font.family': 'DejaVu Sans', 'axes.spines.top': False, 'axes.spines.right': False})
    fig, ax = plt.subplots(figsize=(8, 3.2), facecolor='white')
    ax.set_facecolor('#f8fafc')
    x = np.arange(len(labels))
    ax.bar(x, p, width=0.55, color='#10b981', edgecolor='white', label='Passed', zorder=3)
    ax.bar(x, f, width=0.55, bottom=p, color='#ef4444', edgecolor='white', label='Failed', zorder=3)
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=9, color='#475569', fontweight='bold')
    ax.set_ylabel('Tests', fontsize=8, color='#64748b')
    ax.grid(axis='y', color='#f1f5f9', linewidth=1, zorder=0)
    ax.legend(fontsize=8, frameon=False, loc='upper right')
    fig.text(0.02, 0.97, 'Category Breakdown', fontsize=11, fontweight='bold', color='#1e293b', va='top')
    plt.subplots_adjust(top=0.85, bottom=0.16, left=0.08, right=0.97)
    buf = BIO()
    fig.savefig(buf, format='png', dpi=170, bbox_inches='tight', facecolor='white')
    plt.close()
    buf.seek(0)
    return RLImage(buf, width=155 * mm, height=64 * mm), cats

def _ir_chapter_header(elements, number: str, title: str, color=ACCENT):
    num_display = str(number).zfill(2)
    badge = Table([[Paragraph(
        f'<font color="white" size="13"><b>{num_display}</b></font>',
        ParagraphStyle('IrChapNum', fontSize=13, fontName='Helvetica-Bold', alignment=TA_CENTER, leading=15))
    ]], colWidths=[14 * mm], rowHeights=[14 * mm])
    badge.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), color),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    title_block = Table([
        [Paragraph(f'<font color="#94a3b8" size="7"><b>SECTION {num_display}</b></font>',
                   ParagraphStyle('IrChapEy', fontSize=7, fontName='Helvetica-Bold', leading=8.5))],
        [Paragraph(f'<font color="#1e293b" size="14"><b>{title}</b></font>',
                   ParagraphStyle('IrChapTitle', fontSize=14, fontName='Helvetica-Bold', leading=17))],
    ], colWidths=[148 * mm])
    title_block.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, 0), 0), ('BOTTOMPADDING', (0, 0), (-1, 0), 2),
        ('TOPPADDING', (0, 1), (-1, 1), 0), ('BOTTOMPADDING', (0, 1), (-1, 1), 0),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
    ]))
    row = Table([[badge, title_block]], colWidths=[17 * mm, 151 * mm])
    row.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('LEFTPADDING', (0, 0), (-1, -1), 0)]))
    elements.append(Spacer(1, 10))
    elements.append(row)
    elements.append(HRFlowable(width='100%', thickness=2, color=color, spaceBefore=6, spaceAfter=12))
def _insight_box(text, color, width=161):
    tbl = Table([[Paragraph(
        f'<font color="{color}" size="7.5"><b>AI Analysis: </b></font>'
        f'<font color="#475569" size="7.5">{text}</font>',
        ParagraphStyle('IRInsight', fontSize=7.5, fontName='Helvetica', leading=11))
    ]], colWidths=[width * mm])
    tbl.setStyle(TableStyle([
        ('BACKGROUND',    (0, 0), (-1, -1), HexColor('#f8fafc')),
        ('BOX',           (0, 0), (-1, -1), 0.8, HexColor(color)),
        ('LINEBEFORE',    (0, 0), (0, -1), 3, HexColor(color)),
        ('LEFTPADDING',   (0, 0), (-1, -1), 10),
        ('RIGHTPADDING',  (0, 0), (-1, -1), 10),
        ('TOPPADDING',    (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    return tbl

def _ir_certificate_details(elements, generation_data, tests, score, url, framework, auth_method, grade):
    total_ms = 0
    for t in tests:
        d = str(t.get('duration', '0'))
        try:
            if d.endswith('ms'):
                total_ms += float(d.replace('ms', '') or 0)
            elif d.endswith('s'):
                total_ms += float(d.replace('s', '') or 0) * 1000
        except Exception:
            pass
    exec_time_disp = f'{total_ms/1000:.2f}s' if total_ms else 'N/A'

    details = [
        ('Validation Date',      datetime.now().strftime('%Y-%m-%d  %H:%M'), '#0EA5E9', HexColor('#f0f9ff')),
        ('Framework',            framework, '#0EA5E9', HexColor('#f0f9ff')),
        ('Authentication',       auth_method, '#6366f1', HexColor('#eef2ff')),
        ('Base URL',             url, '#6366f1', HexColor('#eef2ff')),
        ('Execution Time',       exec_time_disp, '#8b5cf6', HexColor('#f5f3ff')),
        ('Total Tests',          str(len(tests)), '#8b5cf6', HexColor('#f5f3ff')),
        ('Overall Grade',        grade, '#10b981', HexColor('#f0fdf4')),
        ('AI Validation Status', 'Verified by NexTest AI', '#10b981', HexColor('#f0fdf4')),
    ]

    rows = []
    row_styles = []
    for i in range(0, len(details), 2):
        chunk = details[i:i + 2]
        row = [Paragraph(
            f'<font color="{color}" size="7.5"><b>{lbl}</b></font><br/>'
            f'<font color="#1e293b" size="9">{val}</font>',
            ParagraphStyle('IRCertDet', fontSize=8, fontName='Helvetica', leading=12))
            for lbl, val, color, bg in chunk]
        if len(row) < 2:
            row.append(Paragraph('', ParagraphStyle('IRCertDetEmpty')))
        rows.append(row)
        ri = len(rows) - 1
        left_color = chunk[0][2]
        row_bg = chunk[0][3]
        row_styles.append(('BACKGROUND', (0, ri), (-1, ri), row_bg))
        row_styles.append(('LINEBEFORE', (0, ri), (0, ri), 3, HexColor(left_color)))
        if len(chunk) > 1:
            row_styles.append(('LINEBEFORE', (1, ri), (1, ri), 1, BORDER_DARK))

    det_tbl = Table(rows, colWidths=[84 * mm, 84 * mm])
    det_tbl.setStyle(TableStyle([
        ('BOX', (0, 0), (-1, -1), 0.8, BORDER_DARK), ('LINEBELOW', (0, 0), (-1, -2), 0.3, BORDER),
        ('LEFTPADDING', (0, 0), (-1, -1), 10), ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ] + row_styles))
    elements.append(KeepTogether([
        Spacer(1, 8),
        Paragraph('<font color="#64748b" size="7.5"><i>Technical details of this validation run — '
                  'environment, timing, and scoring — provided for traceability and audit purposes.'
                  '</i></font>', ParagraphStyle('IRCertDetInfo', fontSize=7.5, fontName='Helvetica', leading=10)),
        Spacer(1, 6),
        det_tbl,
    ]))
    elements.append(Spacer(1, 10))
# ─────────────────────────────────────────────────────────────────────────────
# MAIN — generate_internal_regression_pdf
# ─────────────────────────────────────────────────────────────────────────────

def generate_internal_regression_pdf(generation_data: dict) -> bytes:
    buffer = BytesIO()
    tests, ai_data = _extract_inputs(generation_data)
    url = _base_url(generation_data)
    framework = generation_data.get('framework', 'Playwright')
    auth_method = _auth_method(generation_data)

    pass_count = sum(1 for t in tests if t.get('status') == 'pass')
    fail_count = sum(1 for t in tests if t.get('status') == 'fail')
    skip_count = sum(1 for t in tests if t.get('status') not in ('pass', 'fail'))
    total = len(tests) or 1
    pass_rate = round(pass_count / total * 100)
    rate_color = '#10b981' if pass_rate >= 80 else '#f59e0b' if pass_rate >= 50 else '#ef4444'

    score, score_label, score_color = _compute_score(tests)
    grade, _ = _grade_from_score(score)
    critical_fail = any(_is_critical_fail(t) for t in tests)

    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        rightMargin=20 * mm, leftMargin=22 * mm,
        topMargin=57 * mm, bottomMargin=20 * mm,
    )

    def on_page(canvas, doc):
        W, H = A4
        canvas.saveState()
        canvas.setFillColor(NAVY)
        canvas.rect(0, H - 52 * mm, W, 52 * mm, fill=1, stroke=0)
        canvas.setFillColor(ACCENT)
        canvas.rect(0, H - 54 * mm, W, 2 * mm, fill=1, stroke=0)
        canvas.setFillColor(ACCENT)
        canvas.rect(0, 0, 3, H - 54 * mm, fill=1, stroke=0)
        canvas.setFillColor(LIGHT_BG)
        canvas.rect(0, 0, W, 14 * mm, fill=1, stroke=0)
        canvas.setFillColor(BORDER)
        canvas.rect(0, 14 * mm, W, 0.5, fill=1, stroke=0)
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(MUTED)
        canvas.drawString(20 * mm, 5 * mm, 'Generated by NexTest — Internal Regression Test Report')
        canvas.drawRightString(W - 20 * mm, 5 * mm, f'Page {doc.page}  •  {datetime.now().strftime("%Y-%m-%d")}')
        canvas.restoreState()

    elements = []

    # ── COVER ────────────────────────────────────────────────────────────
    header_tbl = Table([[
        Paragraph('<font color="#eab308"><b>NEX</b></font><font color="#ffffff">TEST</font>',
                  ParagraphStyle('IRLogo', fontSize=24, fontName='Helvetica-Bold')),
        Paragraph(f'<font color="#64748b">Generated</font><br/>'
                  f'<font color="#94a3b8">{datetime.now().strftime("%B %d, %Y  •  %H:%M")}</font>',
                  ParagraphStyle('IRDate', fontSize=8.5, fontName='Helvetica', alignment=TA_RIGHT, leading=13)),
    ]], colWidths=[90 * mm, 78 * mm])
    header_tbl.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(Spacer(1, -38 * mm))
    elements.append(header_tbl)
    elements.append(Spacer(1, 6 * mm))
    elements.append(Paragraph('Internal Regression Test Report',
                               ParagraphStyle('IRTitle', fontSize=22, textColor=WHITE,
                                              fontName='Helvetica-Bold', spaceAfter=2)))
    elements.append(Spacer(1, 14 * mm))

    info_tbl = Table([
        [Paragraph('<font color="#64748b">Base URL</font>', ParagraphStyle('IRIL1', fontSize=8, fontName='Helvetica-Bold', leading=12)),
         Paragraph(f'<font color="#1e293b">{url}</font>', ParagraphStyle('IRIV1', fontSize=8.5, fontName='Helvetica', leading=12))],
        [Paragraph('<font color="#64748b">Framework</font>', ParagraphStyle('IRIL2', fontSize=8, fontName='Helvetica-Bold', leading=12)),
         Paragraph(f'<font color="#f97316"><b>{framework}</b></font>', ParagraphStyle('IRIV2', fontSize=8.5, fontName='Helvetica', leading=12))],
        [Paragraph('<font color="#64748b">Test Type</font>', ParagraphStyle('IRIL3', fontSize=8, fontName='Helvetica-Bold', leading=12)),
         Paragraph('<font color="#eab308"><b>Internal Regression Test</b></font>', ParagraphStyle('IRIV3', fontSize=8.5, fontName='Helvetica', leading=12))],
        [Paragraph('<font color="#64748b">Auth Method</font>', ParagraphStyle('IRIL4', fontSize=8, fontName='Helvetica-Bold', leading=12)),
         Paragraph(f'<font color="#1e293b">{auth_method}</font>', ParagraphStyle('IRIV4', fontSize=8.5, fontName='Helvetica', leading=12))],
        [Paragraph('<font color="#64748b">Generated</font>', ParagraphStyle('IRIL5', fontSize=8, fontName='Helvetica-Bold', leading=12)),
         Paragraph(f'<font color="#1e293b">{datetime.now().strftime("%Y-%m-%d  %H:%M")}</font>', ParagraphStyle('IRIV5', fontSize=8.5, fontName='Helvetica', leading=12))],
    ], colWidths=[32 * mm, 136 * mm])
    info_tbl.setStyle(TableStyle([
        ('BACKGROUND',    (0, 0), (0, -1), LIGHT_BG),
        ('PADDING',       (0, 0), (-1, -1), 7),
        ('LINEBELOW',     (0, 0), (-1, -2), 0.4, BORDER),
        ('BOX',           (0, 0), (-1, -1), 0.8, BORDER_DARK),
        ('ROWBACKGROUNDS', (0, 0), (-1, -1), [WHITE, LIGHT_BG]),
        ('LEFTPADDING',   (0, 0), (0, -1), 10),
        ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(info_tbl)
    elements.append(Spacer(1, 20))

    elements.append(section_header('📊', 'Test Summary', ACCENT))
    elements.append(Spacer(1, 8))
    stats_data = [[
        stat_card(pass_count, 'PASSED', '#10b981', GREEN_BG),
        stat_card(fail_count, 'FAILED', '#ef4444', RED_BG),
        stat_card(skip_count, 'SKIPPED', '#f59e0b', ORANGE_BG),
        stat_card(f'{pass_rate}%', 'PASS RATE', rate_color,
                  GREEN_BG if pass_rate >= 80 else ORANGE_BG if pass_rate >= 50 else RED_BG),
        stat_card(total, 'TOTAL', '#3b82f6', BLUE_BG),
    ]]
    outer = Table(stats_data, colWidths=[33.6 * mm] * 5)
    outer.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(outer)
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(
        '<font color="#94a3b8" size="7"><i>Pass Rate is the raw proportion of tests that succeeded. '
        'The score below is severity-weighted — Authentication and Critical/High severity tests count more.'
        '</i></font>', ParagraphStyle('IRRateNote', fontSize=7, fontName='Helvetica', leading=9)))
    elements.append(Spacer(1, 16))
 

    _ir_chapter_header(elements, '1', 'Regression Test Summary')

    # ── OVERVIEW HERO ────────────────────────────────────────────────────
    overall_status = 'CRITICAL ISSUES' if critical_fail else ('PASSED W/ WARNINGS' if fail_count > 0 else 'ALL CHECKS PASSED')
    status_color = '#ef4444' if critical_fail else ('#f59e0b' if fail_count > 0 else '#10b981')
    risk_level = 'HIGH' if critical_fail else ('MEDIUM' if fail_count > 0 else 'LOW')
    deploy_text = 'NOT READY' if critical_fail else ('READY W/ CAUTION' if fail_count > 0 else 'READY')
    status_explain = (
        'One or more critical checks failed (authentication, or a critical/high severity page). '
        'These block core user journeys — resolve before the next deployment.'
        if critical_fail else
        (f'Core authentication is confirmed, but {fail_count} secondary check(s) failed. '
         'Review the failing pages below.' if fail_count > 0 else
         f'Every regression check executed against {url} passed, including authentication. '
         'This application is stable after the latest changes.')
    )
    elements.append(section_header('🏁', 'Regression Test Overview', ACCENT))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(f'<font color="#64748b" size="7.5"><i>{status_explain}</i></font>',
                               ParagraphStyle('IRStatusExplain', fontSize=7.5, fontName='Helvetica', leading=10)))
    elements.append(Spacer(1, 8))

    def _hero_stat(label, value, color):
        return Table([[Paragraph(
            f'<font color="#94a3b8" size="7"><b>{label}</b></font><br/>'
            f'<font color="{color}" size="12"><b>{value}</b></font>',
            ParagraphStyle('IRHeroStat', fontSize=10, fontName='Helvetica', leading=16, alignment=TA_CENTER))
        ]], colWidths=[40 * mm], style=[
            ('BACKGROUND', (0, 0), (-1, -1), LIGHT_BG),
            ('BOX', (0, 0), (-1, -1), 1, HexColor(color)),
            ('TOPPADDING', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ])
    row = Table([[
        _hero_stat('OVERALL STATUS', overall_status, status_color),
        _hero_stat('REGRESSION SCORE', f'{score}/100', score_color),
        _hero_stat('RISK LEVEL', risk_level, status_color),
        _hero_stat('DEPLOYMENT', deploy_text, status_color),
    ]], colWidths=[42 * mm] * 4)
    row.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(row)
    elements.append(Spacer(1, 10))

    ai_summary = (ai_data or {}).get('summary', '')
    if not ai_summary:
        ai_summary = (
            f'This internal regression run executed {total} test(s) against <b>{url}</b>, with '
            f'{pass_count} passed and {fail_count} failed ({pass_rate}% pass rate). '
            + ('Critical failures were detected on this application.' if critical_fail
               else 'No critical failures were detected on this application.')
        )
    elements.append(_insight_box(ai_summary, '#4f46e5'))
    elements.append(Spacer(1, 16))

    # ── SCORE GAUGE ──────────────────────────────────────────────────────
    elements.append(section_header('🎯', 'Regression Score', HexColor(score_color)))
    elements.append(Spacer(1, 10))
    try:
        gauge_img = _make_score_gauge(score, score_color)
    except Exception:
        gauge_img = None
    right_col = [
        Paragraph(f'<font color="{score_color}" size="15"><b>{score_label}</b></font>',
                  ParagraphStyle('IRHeroLabel', fontSize=15, fontName='Helvetica-Bold', leading=18)),
        Spacer(1, 6),
        Paragraph(
            '<font color="#475569" size="8.5">The Regression Score weighs authentication checks and '
            'critical/high severity pages three times as heavily as low/medium severity checks. '
            'Excellent ≥ 90 · Good ≥ 75 · Acceptable ≥ 50 · Critical below.</font>',
            ParagraphStyle('IRHeroDesc', fontSize=8.5, fontName='Helvetica', leading=13)),
    ]
    gauge_row = Table([[gauge_img, right_col]] if gauge_img else [[right_col]],
                       colWidths=[56 * mm, 112 * mm] if gauge_img else [168 * mm])
    gauge_row.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (0, 0), 'CENTER'),
        ('BOX', (0, 0), (-1, -1), 1.2, HexColor(score_color)),
        ('BACKGROUND', (0, 0), (-1, -1), LIGHT_BG),
        ('LEFTPADDING', (0, 0), (-1, -1), 14),
        ('RIGHTPADDING', (0, 0), (-1, -1), 14),
        ('TOPPADDING', (0, 0), (-1, -1), 14),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 14),
    ]))
    elements.append(gauge_row)
    elements.append(Spacer(1, 18))

    # ── SCENARIOS ────────────────────────────────────────────────────────
    elements.append(section_header('📋', 'Regression Test Scenarios', INDIGO))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(
        f'<font color="#64748b" size="7.5"><i>Regression test plan for <b>{url}</b> — {total} '
        f'scenario(s) executed against the internal application.</i></font>',
        ParagraphStyle('IRScInfo', fontSize=7.5, fontName='Helvetica', leading=10)))
    elements.append(Spacer(1, 8))
    hdr = [
        Paragraph('<font color="#ffffff"><b>#</b></font>', ParagraphStyle('IRSH0', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Test Scenario</b></font>', ParagraphStyle('IRSH1', fontSize=8, fontName='Helvetica-Bold')),
        Paragraph('<font color="#ffffff"><b>Category</b></font>', ParagraphStyle('IRSH2', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Priority</b></font>', ParagraphStyle('IRSH3', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Expected</b></font>', ParagraphStyle('IRSH4', fontSize=8, fontName='Helvetica-Bold')),
    ]
    rows = [hdr]
    row_styles = []
    for i, t in enumerate(tests):
        cat_label = _category_label(t)
        cc = CAT_COLORS.get((t.get('category') or 'navigation').lower(), '#64748b')
        priority = (t.get('priority') or t.get('severity') or 'medium').lower()
        pc = PRIORITY_COLORS.get(priority, '#f59e0b')
        expected = t.get('expected') or t.get('description') or 'Page/element behaves as expected'
        rows.append([
            Paragraph(f'<font color="#64748b"><b>{i+1}</b></font>', ParagraphStyle('IRSID', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<b><font color="#1e293b" size="8">{t.get("name","")}</font></b>', ParagraphStyle('IRSN', fontSize=8, fontName='Helvetica', leading=11)),
            Paragraph(f'<font color="{cc}"><b>{cat_label.upper()}</b></font>', ParagraphStyle('IRSC', fontSize=7, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="{pc}"><b>{priority.upper()}</b></font>', ParagraphStyle('IRSP', fontSize=7, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="#475569" size="7">{str(expected)[:75]}</font>', ParagraphStyle('IRSE', fontSize=7, fontName='Helvetica', leading=10)),
        ])
        if i % 2 == 1:
            row_styles.append(('BACKGROUND', (0, i + 1), (-1, i + 1), LIGHT_BG))
    tbl = Table(rows, colWidths=[8 * mm, 56 * mm, 26 * mm, 20 * mm, 58 * mm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), NAVY),
        ('PADDING', (0, 0), (-1, -1), 7),
        ('LINEBELOW', (0, 0), (-1, -1), 0.4, BORDER),
        ('BOX', (0, 0), (-1, -1), 0.8, INDIGO),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (3, -1), 'CENTER'),
        ('LINEBEFORE', (2, 1), (2, -1), 1, BORDER),
        ('LINEBEFORE', (4, 1), (4, -1), 1, BORDER),
    ] + row_styles))
    elements.append(tbl)
    elements.append(Spacer(1, 16))

    # ── RESULTS BY CATEGORY ──────────────────────────────────────────────
    elements.append(section_header('📊', 'Results by Category', ACCENT))
    elements.append(Spacer(1, 8))
    cats = {}
    for t in tests:
        c = _category_label(t)
        if c not in cats:
            cats[c] = {'pass': 0, 'fail': 0, 'total': 0}
        cats[c]['total'] += 1
        if t.get('status') == 'pass':
            cats[c]['pass'] += 1
        elif t.get('status') == 'fail':
            cats[c]['fail'] += 1
    hdr2 = [
        Paragraph('<font color="#ffffff"><b>Category</b></font>', ParagraphStyle('IRCH1', fontSize=8, fontName='Helvetica-Bold')),
        Paragraph('<font color="#ffffff"><b>Total</b></font>', ParagraphStyle('IRCH2', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Passed</b></font>', ParagraphStyle('IRCH3', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Failed</b></font>', ParagraphStyle('IRCH4', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Pass Rate</b></font>', ParagraphStyle('IRCH5', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Verdict</b></font>', ParagraphStyle('IRCH6', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
    ]
    rows2 = [hdr2]
    row_styles2 = []
    for cat in CAT_ORDER:
        d = cats.get(cat)
        if not d or d['total'] == 0:
            continue
        rate = round(d['pass'] / d['total'] * 100)
        rate_c = '#10b981' if rate == 100 else '#f59e0b' if rate >= 60 else '#ef4444'
        verdict = 'PASS' if d['fail'] == 0 else 'FAIL'
        vc = '#10b981' if d['fail'] == 0 else '#ef4444'
        row_bg = HexColor('#f0fdf4') if d['fail'] == 0 else HexColor('#fef2f2')
        rows2.append([
            Paragraph(f'<font color="{CAT_COLORS.get(cat.lower(),"#64748b")}"><b>{cat}</b></font>', ParagraphStyle('IRCC', fontSize=8, fontName='Helvetica-Bold')),
            Paragraph(f'<font color="#1e293b"><b>{d["total"]}</b></font>', ParagraphStyle('IRCT', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="#10b981"><b>{d["pass"]}</b></font>', ParagraphStyle('IRCP', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="#ef4444"><b>{d["fail"]}</b></font>', ParagraphStyle('IRCF', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="{rate_c}"><b>{rate}%</b></font>', ParagraphStyle('IRCR', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="{vc}"><b>{verdict}</b></font>', ParagraphStyle('IRCV', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        ])
        ri = len(rows2) - 1
        row_styles2.append(('BACKGROUND', (0, ri), (-1, ri), row_bg))
    tbl2 = Table(rows2, colWidths=[36 * mm, 22 * mm, 22 * mm, 22 * mm, 26 * mm, 40 * mm], repeatRows=1)
    tbl2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), NAVY),
        ('PADDING', (0, 0), (-1, -1), 8),
        ('LINEBELOW', (0, 0), (-1, -1), 0.4, BORDER),
        ('BOX', (0, 0), (-1, -1), 0.8, ACCENT),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (1, 0), (5, -1), 'CENTER'),
    ] + row_styles2))
    elements.append(tbl2)
    elements.append(Spacer(1, 16))

    # ── CHART ────────────────────────────────────────────────────────────
    _ir_chapter_header(elements, '2', 'Regression Coverage')
    try:
        chart_img, cats_data = _make_category_chart(tests)
        worst_cat = max(cats_data, key=lambda c: cats_data[c]['fail']) if cats_data else None
        cat_insight = (
            f'{worst_cat} currently has the most failures ({cats_data[worst_cat]["fail"]}) — '
            f'this is the category to prioritize first.'
            if worst_cat and cats_data[worst_cat]['fail'] > 0
            else 'No category shows any failures — coverage is currently clean across the board.'
        )
        framed = Table([[chart_img]], colWidths=[161 * mm])
        framed.setStyle(TableStyle([
            ('BOX', (0, 0), (-1, -1), 1, ACCENT),
            ('BACKGROUND', (0, 0), (-1, -1), WHITE),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        elements.append(section_header('📈', 'Category Breakdown Chart', ACCENT))
        elements.append(Spacer(1, 8))
        elements.append(framed)
        elements.append(Spacer(1, 6))
        elements.append(_insight_box(cat_insight, '#0EA5E9'))
        elements.append(Spacer(1, 16))
    except Exception as e:
        print(f"[INTERNAL-REGRESSION PDF] chart error: {e}")

    # ── ENVIRONMENT ──────────────────────────────────────────────────────
    _ir_chapter_header(elements, '3', 'Test Details')

    elements.append(section_header('⚙️', 'Execution Environment', ACCENT))
    elements.append(Spacer(1, 6))
    elements.append(Paragraph(
        '<font color="#64748b" size="7.5"><i>Framework and authentication method used to run this '
        'internal regression audit — for reproducibility of the results above.</i></font>',
        ParagraphStyle('IREnvInfo', fontSize=7.5, fontName='Helvetica', leading=10)))
    elements.append(Spacer(1, 8))

    total_ms = 0
    for t in tests:
        d = str(t.get('duration', '0'))
        try:
            if d.endswith('ms'):
                total_ms += float(d.replace('ms', '') or 0)
            elif d.endswith('s'):
                total_ms += float(d.replace('s', '') or 0) * 1000
        except Exception:
            pass

    env_items = [
        ('FRAMEWORK', framework, '#0EA5E9'),
        ('AUTHENTICATION METHOD', auth_method, '#6366f1'),
        ('BASE URL', url, '#f59e0b'),
        ('TOTAL TESTS', str(total), '#10b981'),
        ('EXECUTION TIME', f'{total_ms/1000:.2f}s' if total_ms else 'N/A', '#ec4899'),
        ('NEXTEST VERSION', generation_data.get('nextest_version', '1.0.0'), '#8b5cf6'),
    ]

    def _env_card(label, value, color):
        cell = Table([[Paragraph(
            f'<font color="{color}" size="7"><b>{label}</b></font><br/>'
            f'<font color="#1e293b" size="10"><b>{value}</b></font>',
            ParagraphStyle('IREnvC', fontSize=9, fontName='Helvetica', leading=15, alignment=TA_CENTER))
        ]], colWidths=[54 * mm])
        cell.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), WHITE),
            ('BOX', (0, 0), (-1, -1), 1, HexColor(color)),
            ('LINEABOVE', (0, 0), (-1, 0), 3, HexColor(color)),
            ('TOPPADDING', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
        return cell

    card_rows = []
    for i in range(0, len(env_items), 3):
        chunk = env_items[i:i + 3]
        cells = [_env_card(l, v, c) for l, v, c in chunk]
        while len(cells) < 3:
            cells.append(Paragraph('', ParagraphStyle('IREnvEmpty')))
        card_rows.append(cells)
    outer2 = Table(card_rows, colWidths=[56 * mm] * 3)
    outer2.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(outer2)
    elements.append(Spacer(1, 16))

    # ── DETAILED RESULTS ─────────────────────────────────────────────────
    elements.append(section_header('🧪', 'Detailed Test Results', TEAL))
    elements.append(Paragraph(
        '<font color="#64748b" size="7.5"><i>Real results from Playwright execution. '
        'Every value comes directly from the test runner.</i></font>',
        ParagraphStyle('IRDRInfo', fontSize=7.5, fontName='Helvetica', leading=10)))
    elements.append(Spacer(1, 8))

    hdr3 = [
        Paragraph('<font color="#ffffff"><b>#</b></font>', ParagraphStyle('IRDH0', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Test Name</b></font>', ParagraphStyle('IRDH1', fontSize=8, fontName='Helvetica-Bold')),
        Paragraph('<font color="#ffffff"><b>Category</b></font>', ParagraphStyle('IRDH2', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Status</b></font>', ParagraphStyle('IRDH3', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        Paragraph('<font color="#ffffff"><b>Result / Reason</b></font>', ParagraphStyle('IRDH4', fontSize=8, fontName='Helvetica-Bold')),
        Paragraph('<font color="#ffffff"><b>Duration</b></font>', ParagraphStyle('IRDH5', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
    ]
    rows3 = [hdr3]
    row_styles3 = []
    for i, t in enumerate(tests):
        status = t.get('status', 'skip')
        sc = '#10b981' if status == 'pass' else '#ef4444' if status == 'fail' else '#f59e0b'
        s_label = '✓  PASS' if status == 'pass' else '✗  FAIL' if status == 'fail' else '■  SKIP'
        s_bg = HexColor('#f0fdf4') if status == 'pass' else HexColor('#fef2f2') if status == 'fail' else HexColor('#fffbeb')
        cat_label = _category_label(t)
        cc = CAT_COLORS.get((t.get('category') or '').lower(), '#64748b')
        reason = t.get('reason') or t.get('reason_pass') or '—'
        duration = t.get('duration', '—')
        rows3.append([
            Paragraph(f'<font color="#64748b"><b>{i+1}</b></font>', ParagraphStyle('IRDID', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<b><font color="#1e293b" size="8">{t.get("name","")}</font></b>', ParagraphStyle('IRDN', fontSize=8, fontName='Helvetica', leading=11)),
            Paragraph(f'<font color="{cc}"><b>{cat_label.upper()}</b></font>', ParagraphStyle('IRDC', fontSize=7, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="{sc}"><b>{s_label}</b></font>', ParagraphStyle('IRDS', fontSize=7.5, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph(f'<font color="#475569" size="7">{str(reason)[:90]}</font>', ParagraphStyle('IRDR', fontSize=7, fontName='Helvetica', leading=10)),
            Paragraph(f'<font color="#64748b" size="7"><b>{duration}</b></font>', ParagraphStyle('IRDD', fontSize=7, fontName='Helvetica-Bold', alignment=TA_CENTER)),
        ])
        row_styles3.append(('BACKGROUND', (3, i + 1), (3, i + 1), s_bg))
    tbl3 = Table(rows3, colWidths=[8 * mm, 46 * mm, 24 * mm, 20 * mm, 52 * mm, 18 * mm], repeatRows=1)
    tbl3.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), NAVY),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [WHITE, LIGHT_BG]),
        ('PADDING', (0, 0), (-1, -1), 7),
        ('LINEBELOW', (0, 0), (-1, -1), 0.4, BORDER),
        ('BOX', (0, 0), (-1, -1), 0.8, TEAL),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (3, -1), 'CENTER'),
        ('ALIGN', (5, 0), (5, -1), 'CENTER'),
        ('LINEBEFORE', (4, 1), (4, -1), 1, BORDER),
    ] + row_styles3))
    elements.append(tbl3)
    elements.append(Spacer(1, 16))

    # ── AI RECOMMENDATIONS ───────────────────────────────────────────────
    _ir_chapter_header(elements, '4', 'AI Insights & Recommendations')
    recs = (ai_data or {}).get('recommendations', []) or []
    elements.append(section_header('🤖', 'AI Recommendations', HexColor('#4f46e5')))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(
        '<font color="#64748b" size="7.5"><i>Recommendations generated only for issues detected '
        'on this application.</i></font>', ParagraphStyle('IRRSInfo', fontSize=7.5, fontName='Helvetica', leading=10)))
    elements.append(Spacer(1, 8))
    if recs:
        hdr4 = [
            Paragraph('<font color="#ffffff"><b>Priority</b></font>', ParagraphStyle('IRRH1', fontSize=8, fontName='Helvetica-Bold', alignment=TA_CENTER)),
            Paragraph('<font color="#ffffff"><b>Category</b></font>', ParagraphStyle('IRRH2', fontSize=8, fontName='Helvetica-Bold')),
            Paragraph('<font color="#ffffff"><b>Issue</b></font>', ParagraphStyle('IRRH3', fontSize=8, fontName='Helvetica-Bold')),
            Paragraph('<font color="#ffffff"><b>Recommended Fix</b></font>', ParagraphStyle('IRRH4', fontSize=8, fontName='Helvetica-Bold')),
        ]
        rows4 = [hdr4]
        row_styles4 = []
        for i, rec in enumerate(recs):
            priority = (rec.get('priority') or 'medium').lower()
            pc = PRIORITY_COLORS.get(priority, '#f59e0b')
            rows4.append([
                Paragraph(f'<font color="{pc}"><b>{priority.upper()}</b></font>', ParagraphStyle('IRRP', fontSize=7.5, fontName='Helvetica-Bold', alignment=TA_CENTER)),
                Paragraph(f'<font color="#4f46e5" size="7.5"><b>{(rec.get("category","") or "").upper()}</b></font>', ParagraphStyle('IRRC', fontSize=7.5, fontName='Helvetica-Bold')),
                Paragraph(f'<font color="#1e293b" size="7.5">{rec.get("issue","")[:70]}</font>', ParagraphStyle('IRRI', fontSize=7.5, fontName='Helvetica', leading=10)),
                Paragraph(f'<font color="#475569" size="7.5">{rec.get("fix","")[:80]}</font>', ParagraphStyle('IRRF', fontSize=7.5, fontName='Helvetica', leading=10)),
            ])
            if i % 2 == 1:
                row_styles4.append(('BACKGROUND', (0, i + 1), (-1, i + 1), LIGHT_BG))
        tbl4 = Table(rows4, colWidths=[20 * mm, 26 * mm, 56 * mm, 66 * mm], repeatRows=1)
        tbl4.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('PADDING', (0, 0), (-1, -1), 7),
            ('LINEBELOW', (0, 0), (-1, -1), 0.4, BORDER),
            ('BOX', (0, 0), (-1, -1), 0.8, HexColor('#4f46e5')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ] + row_styles4))
        elements.append(tbl4)
    else:
        elements.append(Paragraph(
            '<font color="#94a3b8" size="8">No specific issues flagged by AI for this run.</font>',
            ParagraphStyle('IRNoRec', fontSize=8, fontName='Helvetica')))
    elements.append(Spacer(1, 16))

    # ── ACTION PLAN ──────────────────────────────────────────────────────
    action_plan = (ai_data or {}).get('action_plan', []) or []
    if action_plan:
        elements.append(section_header('📋', 'AI-Generated Action Plan', GOLD))
        elements.append(Spacer(1, 8))
        for i, step in enumerate(action_plan):
            step_text = step if isinstance(step, str) else step.get('action', '')
            elements.append(Paragraph(
                f'<font color="#64748b" size="8.5"><b>Step {i+1}:</b> {step_text}</font>',
                ParagraphStyle('IRAPStep', fontSize=8.5, fontName='Helvetica', leading=13)))
        elements.append(Spacer(1, 16))

    # ── EXECUTIVE SUMMARY ────────────────────────────────────────────────
    if recs:
        PRI_ORDER = {'high': 0, 'medium': 1, 'low': 2}
        top = sorted(recs, key=lambda x: PRI_ORDER.get((x.get('priority') or 'medium').lower(), 1))[:2]
        elements.append(section_header('⭐', 'Executive Summary', GOLD))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph('<font color="#1e293b" size="9"><b>Top Priority Actions</b></font>',
                                   ParagraphStyle('IRExecH', fontSize=9, fontName='Helvetica-Bold')))
        elements.append(Spacer(1, 8))
        for i, item in enumerate(top):
            card = Table([[Paragraph(
                f'<font color="#c9a227" size="11"><b>{i+1}</b></font><br/>'
                f'<font color="#4f46e5" size="8"><b>{(item.get("category","") or "").upper()}</b></font><br/>'
                f'<font color="#1e293b" size="8.5"><b>{item.get("issue","")}</b></font><br/>'
                f'<font color="#64748b" size="7.5">{item.get("fix","")}</font>',
                ParagraphStyle('IRExecCard', fontSize=8.5, fontName='Helvetica', leading=12))
            ]], colWidths=[168 * mm])
            card.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), LIGHT_BG),
                ('BOX', (0, 0), (-1, -1), 0.8, GOLD),
                ('LEFTPADDING', (0, 0), (-1, -1), 14),
                ('TOPPADDING', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ]))
            elements.append(card)
            elements.append(Spacer(1, 8))
        insight = (
            f'These {len(top)} action(s) target the largest contributors to the current score '
            f'of {score}/100. Re-run the regression suite after applying them to confirm improvement.'
        )
        elements.append(_insight_box(insight, '#c9a227'))
        elements.append(Spacer(1, 8))

    

     # ── FINAL VERDICT ────────────────────────────────────────────────────
    _ir_chapter_header(elements, '5', 'Report Conclusion')
    elements.append(section_header('🏁', 'Final AI Verdict', ACCENT))
    elements.append(Spacer(1, 6))
    if critical_fail:
        vc, vb, vi = '#ef4444', HexColor('#fef2f2'), '🔴'
        vt = (f'Internal Regression Test FAILED — critical checks did not pass on {url}. '
              f'Authentication or high-severity page issues were detected. This application is '
              f'NOT ready for the next deployment until resolved.')
    elif fail_count > 0:
        vc, vb, vi = '#b45309', HexColor('#fffbeb'), '🟡'
        vt = (f'Internal Regression Test passed with {fail_count} non-critical issue(s) on {url}. '
              f'The application is accessible and stable, but the failing checks should be '
              f'reviewed before the next deployment.')
    else:
        vc, vb, vi = '#059669', HexColor('#f0fdf4'), '🟢'
        vt = (f'Internal Regression Test PASSED — all {pass_count} test(s) succeeded on {url}. '
              f'This application is stable after the latest changes.')

    final_tbl = Table([[Paragraph(
        f'<font color="{vc}" size="9"><b>{vi}  Final Regression Verdict</b></font><br/>'
        f'<font color="{vc}" size="8">{vt}</font><br/><br/>'
        f'<font color="#64748b" size="8"><b>Regression Score: </b></font>'
        f'<font color="{vc}" size="8"><b>{score}/100</b></font>'
        f'<font color="#94a3b8" size="8">    |    </font>'
        f'<font color="#64748b" size="8"><b>Risk Level: </b></font>'
        f'<font color="{vc}" size="8"><b>{risk_level}</b></font>'
        f'<font color="#94a3b8" size="8">    |    </font>'
        f'<font color="#64748b" size="8"><b>Deployment Readiness: </b></font>'
        f'<font color="{vc}" size="8"><b>{score_label}</b></font>',
        ParagraphStyle('IRFV', fontSize=8, fontName='Helvetica', leading=13))
    ]], colWidths=[168 * mm])
    final_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), vb),
        ('BOX', (0, 0), (-1, -1), 2, HexColor(vc)),
        ('LEFTPADDING', (0, 0), (-1, -1), 14),
        ('RIGHTPADDING', (0, 0), (-1, -1), 14),
        ('TOPPADDING', (0, 0), (-1, -1), 12),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
    ]))
    elements.append(final_tbl)
    elements.append(Spacer(1, 20))

    # ── CERTIFICATE ──────────────────────────────────────────────────────
    elements.append(section_header('🏆', 'Certificate of Internal Regression Validation', ACCENT))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(
        '<font color="#64748b" size="7.5"><i>Official validation summary confirming the outcome '
        'of this internal regression run — issued automatically by NexTest AI.</i></font>',
        ParagraphStyle('IRCertInfo', fontSize=7.5, fontName='Helvetica', leading=10)))
    elements.append(Spacer(1, 8))
    cert_data = {'global_score': score, 'score_label': score_label, 'score_color': score_color}
    elements.append(_build_certificate_card(cert_data, url, title="CERTIFICATE OF INTERNAL REGRESSION VALIDATION"))
    _ir_certificate_details(elements, generation_data, tests, score, url, framework, auth_method, grade)

    doc.build(elements, onFirstPage=on_page, onLaterPages=on_page)
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# XLSX — generate_internal_regression_xlsx
# ─────────────────────────────────────────────────────────────────────────────

_XLSX_FONT = "Calibri"
_XLSX_THIN = Side(style="thin", color="E2E8F0")
_XLSX_BORDER = Border(left=_XLSX_THIN, right=_XLSX_THIN, top=_XLSX_THIN, bottom=_XLSX_THIN)


def _xh(ws, row, headers, col_start=1):
    for i, h in enumerate(headers, start=col_start):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=_XLSX_FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="0A0F1E")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _XLSX_BORDER


def _xautofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _xtitle(ws, title: str, color_hex: str):
    ws.insert_rows(1, 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name=_XLSX_FONT, bold=True, size=13, color=color_hex)
    c.fill = PatternFill("solid", fgColor="0A0F1E")
    for col in range(1, 7):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor="0A0F1E")
    ws.row_dimensions[1].height = 24


def generate_internal_regression_xlsx(generation_data: dict) -> bytes:
    tests, ai_data = _extract_inputs(generation_data)
    url = _base_url(generation_data)
    framework = generation_data.get('framework', 'Playwright')
    auth_method = _auth_method(generation_data)

    pass_count = sum(1 for t in tests if t.get('status') == 'pass')
    fail_count = sum(1 for t in tests if t.get('status') == 'fail')
    skip_count = sum(1 for t in tests if t.get('status') not in ('pass', 'fail'))
    total = len(tests) or 1
    pass_rate = round(pass_count / total * 100)
    rate_color = "10B981" if pass_rate >= 80 else "F59E0B" if pass_rate >= 50 else "EF4444"

    score, score_label, score_color_hex = _compute_score(tests)
    score_color = score_color_hex.lstrip('#').upper()
    grade, _ = _grade_from_score(score)
    critical_fail = any(_is_critical_fail(t) for t in tests)

    wb = Workbook()

    # ── SHEET 1 — Overview ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H3")
    for r_ in range(1, 4):
        for c_ in range(1, 9):
            ws.cell(row=r_, column=c_).fill = PatternFill("solid", fgColor="0A0F1E")
    ws["A1"] = "NEXTEST — Internal Regression Test Report"
    ws["A1"].font = Font(name=_XLSX_FONT, bold=True, size=20, color="EAB308")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells("A4:H4")
    ws["A4"] = f"Generated {datetime.now().strftime('%B %d, %Y  •  %H:%M')}"
    ws["A4"].font = Font(name=_XLSX_FONT, italic=True, size=9, color="64748B")

    info_rows = [
        ("Base URL", url),
        ("Framework", framework),
        ("Auth Method", auth_method),
        ("Test Type", "Internal Regression Test"),
        ("Regression Score", f"{score}/100 — {score_label}"),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    r = 6
    for lbl, val in info_rows:
        lc = ws.cell(row=r, column=1, value=lbl)
        lc.font = Font(name=_XLSX_FONT, bold=True, color="64748B", size=9)
        lc.fill = PatternFill("solid", fgColor="F8FAFC")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        vc = ws.cell(row=r, column=2, value=val)
        vc.font = Font(name=_XLSX_FONT, size=10, color="1E293B", bold=(lbl == "Regression Score"))
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = _XLSX_BORDER
        r += 1

    r += 2
    stats = [
        ("PASSED", pass_count, "10B981", "D1FAE5"),
        ("FAILED", fail_count, "EF4444", "FEE2E2"),
        ("SKIPPED", skip_count, "F59E0B", "FEF3C7"),
        ("PASS RATE", f"{pass_rate}%", rate_color,
         "D1FAE5" if pass_rate >= 80 else "FEF3C7" if pass_rate >= 50 else "FEE2E2"),
        ("TOTAL", total, "3B82F6", "DBEAFE"),
    ]
    col = 1
    for lbl, val, color, bg in stats:
        c1 = ws.cell(row=r, column=col, value=val)
        c1.font = Font(name=_XLSX_FONT, bold=True, size=16, color=color)
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.fill = PatternFill("solid", fgColor=bg)
        c1.border = _XLSX_BORDER
        c2 = ws.cell(row=r + 1, column=col, value=lbl)
        c2.font = Font(name=_XLSX_FONT, bold=True, size=8, color="64748B")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.fill = PatternFill("solid", fgColor=bg)
        c2.border = _XLSX_BORDER
        col += 1
    ws.row_dimensions[r].height = 22
    ws.row_dimensions[r + 1].height = 18
    r += 3

    overall_status = 'CRITICAL ISSUES' if critical_fail else ('PASSED W/ WARNINGS' if fail_count > 0 else 'ALL CHECKS PASSED')
    risk_level = 'HIGH' if critical_fail else ('MEDIUM' if fail_count > 0 else 'LOW')
    deploy_text = 'NOT READY' if critical_fail else ('READY W/ CAUTION' if fail_count > 0 else 'READY')
    status_color = "EF4444" if critical_fail else ("F59E0B" if fail_count > 0 else "10B981")

    ws.cell(row=r, column=1, value="Regression Test Overview").font = Font(
        name=_XLSX_FONT, bold=True, size=12, color="1E293B")
    r += 1
    hero_items = [
        ("OVERALL STATUS", overall_status, status_color),
        ("REGRESSION SCORE", f"{score}/100", score_color),
        ("RISK LEVEL", risk_level, status_color),
        ("DEPLOYMENT", deploy_text, status_color),
    ]
    col = 1
    for lbl, val, color in hero_items:
        c1 = ws.cell(row=r, column=col, value=val)
        c1.font = Font(name=_XLSX_FONT, bold=True, size=12, color=color)
        c1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c1.border = _XLSX_BORDER
        c2 = ws.cell(row=r + 1, column=col, value=lbl)
        c2.font = Font(name=_XLSX_FONT, bold=True, size=7.5, color="64748B")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.border = _XLSX_BORDER
        col += 1
    ws.row_dimensions[r].height = 26
    ws.row_dimensions[r + 1].height = 16
    r += 3

    ai_summary_text = (ai_data or {}).get('summary', '')
    if not ai_summary_text:
        ai_summary_text = (
            f'This internal regression run executed {total} test(s) against {url}, with {pass_count} '
            f'passed and {fail_count} failed ({pass_rate}% pass rate). '
            + ('Critical failures were detected on this application.' if critical_fail
               else 'No critical failures were detected on this application.'))
    n_rows = max(2, math.ceil(len(ai_summary_text) / 95) + 1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r + n_rows - 1, end_column=8)
    c = ws.cell(row=r, column=1, value=ai_summary_text)
    c.font = Font(name=_XLSX_FONT, size=9, color="475569")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.fill = PatternFill("solid", fgColor="F8FAFC")
    r += n_rows + 1

    # ── Category Breakdown chart ──
    cat_hdr_row = r
    _xh(ws, r, ["Category", "Total", "Passed", "Failed"])
    r += 1
    cats = {c: {'pass': 0, 'fail': 0, 'total': 0} for c in CAT_ORDER}
    for t in tests:
        label = _category_label(t)
        if label not in cats:
            cats[label] = {'pass': 0, 'fail': 0, 'total': 0}
        cats[label]['total'] += 1
        if t.get('status') == 'pass':
            cats[label]['pass'] += 1
        elif t.get('status') == 'fail':
            cats[label]['fail'] += 1
    for cat in CAT_ORDER:
        d = cats.get(cat)
        if not d or d['total'] == 0:
            continue
        cc = CAT_COLORS.get(cat.lower(), '#64748b').lstrip('#').upper()
        c1 = ws.cell(row=r, column=1, value=cat)
        c1.font = Font(name=_XLSX_FONT, bold=True, size=9, color=cc)
        c1.border = _XLSX_BORDER
        for i, v in enumerate([d['total'], d['pass'], d['fail']], start=2):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=_XLSX_FONT, size=9, color="1E293B")
            c.alignment = Alignment(horizontal="center")
            c.border = _XLSX_BORDER
        r += 1
    chart_end = r - 1

    if chart_end >= cat_hdr_row + 1:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Category Breakdown"
        chart.y_axis.title = "Tests"
        cats_ref = Reference(ws, min_col=1, min_row=cat_hdr_row + 1, max_row=chart_end)
        pass_ref = Reference(ws, min_col=3, min_row=cat_hdr_row, max_row=chart_end)
        fail_ref = Reference(ws, min_col=4, min_row=cat_hdr_row, max_row=chart_end)
        chart.add_data(pass_ref, titles_from_data=True)
        chart.add_data(fail_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.series[0].graphicalProperties.solidFill = "10B981"
        chart.series[1].graphicalProperties.solidFill = "EF4444"
        chart.width = 18
        chart.height = 9
        ws.add_chart(chart, f"A{chart_end + 2}")

    _xautofit(ws, [22, 12, 12, 12, 14, 14, 14, 14])
    ws.freeze_panes = "A6"

    # ── SHEET 2 — Test Scenarios ─────────────────────────────────────────
    ws2 = wb.create_sheet("Test Scenarios")
    ws2.sheet_view.showGridLines = False
    _xh(ws2, 1, ["#", "Scenario", "Category", "Priority", "Expected"])
    ws2.freeze_panes = "A2"
    for i, t in enumerate(tests, start=2):
        cat_label = _category_label(t)
        cc = CAT_COLORS.get((t.get('category') or 'navigation').lower(), '#64748b').lstrip('#').upper()
        priority = (t.get('priority') or t.get('severity') or 'medium').lower()
        pc = PRIORITY_COLORS.get(priority, '#f59e0b').lstrip('#').upper()
        expected = t.get('expected') or t.get('description') or 'Page/element behaves as expected'
        vals = [i - 1, t.get('name', ''), cat_label.upper(), priority.upper(), str(expected)[:120]]
        for col, v in enumerate(vals, start=1):
            c = ws2.cell(row=i, column=col, value=v)
            c.font = Font(name=_XLSX_FONT, size=9, bold=(col in (3, 4)),
                          color=(cc if col == 3 else pc if col == 4 else "1E293B"))
            c.alignment = Alignment(horizontal="center" if col in (1, 3, 4) else "left", vertical="top", wrap_text=(col == 5))
            c.border = _XLSX_BORDER
    _xtitle(ws2, "REGRESSION TEST SCENARIOS", "EAB308")
    _xautofit(ws2, [5, 40, 18, 12, 55])

    # ── SHEET 3 — Results by Category ────────────────────────────────────
    ws3 = wb.create_sheet("Results by Category")
    ws3.sheet_view.showGridLines = False
    _xh(ws3, 1, ["Category", "Total", "Passed", "Failed", "Pass Rate", "Verdict"])
    ws3.freeze_panes = "A2"
    row = 2
    for cat in CAT_ORDER:
        d = cats.get(cat)
        if not d or d['total'] == 0:
            continue
        rate = round(d['pass'] / d['total'] * 100)
        verdict = 'PASS' if d['fail'] == 0 else 'FAIL'
        bg = "D1FAE5" if d['fail'] == 0 else "FEE2E2"
        cc = CAT_COLORS.get(cat.lower(), '#64748b').lstrip('#').upper()
        vals = [cat, d['total'], d['pass'], d['fail'], f"{rate}%", verdict]
        for col, v in enumerate(vals, start=1):
            c = ws3.cell(row=row, column=col, value=v)
            c.font = Font(name=_XLSX_FONT, size=9, bold=(col in (1, 6)),
                          color=(cc if col == 1 else ("EF4444" if verdict == "FAIL" else "10B981") if col == 6 else "1E293B"))
            c.fill = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(horizontal="center" if col > 1 else "left", vertical="center")
            c.border = _XLSX_BORDER
        row += 1
    _xtitle(ws3, "RESULTS BY CATEGORY", "EAB308")
    _xautofit(ws3, [22, 10, 10, 10, 12, 12])

    # ── SHEET 4 — Detailed Results ────────────────────────────────────────
    ws4 = wb.create_sheet("Detailed Results")
    ws4.sheet_view.showGridLines = False
    _xh(ws4, 1, ["#", "Test Name", "Category", "Status", "Result / Reason", "Duration"])
    ws4.freeze_panes = "A2"
    STATUS_COLOR = {"pass": ("10B981", "D1FAE5"), "fail": ("EF4444", "FEE2E2"), "skip": ("F59E0B", "FFFBEB")}
    for i, t in enumerate(tests, start=2):
        status = t.get('status', 'skip')
        s_color, s_bg = STATUS_COLOR.get(status, ("94A3B8", "FFFFFF"))
        cat_label = _category_label(t)
        cc = CAT_COLORS.get((t.get('category') or '').lower(), '#64748b').lstrip('#').upper()
        reason = t.get('reason') or t.get('reason_pass') or '—'
        vals = [i - 1, t.get('name', ''), cat_label.upper(), status.upper(), str(reason)[:150], t.get('duration', '—')]
        for col, v in enumerate(vals, start=1):
            c = ws4.cell(row=i, column=col, value=v)
            c.font = Font(name=_XLSX_FONT, size=9, bold=(col in (3, 4)),
                          color=(cc if col == 3 else s_color if col == 4 else "1E293B"))
            c.alignment = Alignment(horizontal="center" if col in (1, 3, 4, 6) else "left", vertical="top", wrap_text=(col == 5))
            c.border = _XLSX_BORDER
            if col == 4:
                c.fill = PatternFill("solid", fgColor=s_bg)
        ws4.row_dimensions[i].height = 20
    _xtitle(ws4, "DETAILED REGRESSION TEST RESULTS", "0D9488")
    _xautofit(ws4, [5, 40, 18, 12, 55, 12])

    # ── SHEET 5 — AI Recommendations ─────────────────────────────────────
    ws5 = wb.create_sheet("AI Recommendations")
    ws5.sheet_view.showGridLines = False
    _xh(ws5, 1, ["Priority", "Category", "Issue", "Fix"])
    recs = (ai_data or {}).get('recommendations', []) or []
    PRI_BG = {"high": "FEE2E2", "medium": "FFFBEB", "low": "D1FAE5"}
    PRI_COLOR = {"high": "EF4444", "medium": "F59E0B", "low": "10B981"}
    r = 2
    for rec in recs:
        pri = (rec.get('priority') or 'medium').lower()
        vals = [pri.upper(), (rec.get('category', '—') or '—').upper(), rec.get('issue', ''),
                rec.get('fix', '')]
        for col, v in enumerate(vals, start=1):
            c = ws5.cell(row=r, column=col, value=v)
            c.font = Font(name=_XLSX_FONT, size=9, bold=(col == 1),
                          color=PRI_COLOR.get(pri, "F59E0B") if col == 1 else "1E293B")
            c.alignment = Alignment(vertical="top", wrap_text=(col >= 3), horizontal="center" if col == 1 else "left")
            c.fill = PatternFill("solid", fgColor=PRI_BG.get(pri, "FFFBEB"))
            c.border = _XLSX_BORDER
        ws5.row_dimensions[r].height = 30
        r += 1
    if not recs:
        ws5.merge_cells("A2:D2")
        c = ws5.cell(row=2, column=1, value="No issues detected on this application — no recommendations needed for this run.")
        c.font = Font(name=_XLSX_FONT, italic=True, size=9, color="94A3B8")
    _xtitle(ws5, "AI RECOMMENDATIONS", "4F46E5")
    _xautofit(ws5, [12, 18, 40, 55])

    # ── SHEET 6 — Action Plan ─────────────────────────────────────────────
    ws6 = wb.create_sheet("Action Plan")
    ws6.sheet_view.showGridLines = False
    _xh(ws6, 1, ["Step", "Action"])
    ws6.freeze_panes = "A2"
    action_plan = (ai_data or {}).get('action_plan', []) or []
    r = 2
    for i, step in enumerate(action_plan, start=1):
        step_text = step if isinstance(step, str) else step.get('action', '')
        c1 = ws6.cell(row=r, column=1, value=f"Step {i}")
        c1.font = Font(name=_XLSX_FONT, bold=True, size=9, color="C9A227")
        c1.alignment = Alignment(vertical="top")
        c1.border = _XLSX_BORDER
        c2 = ws6.cell(row=r, column=2, value=step_text)
        c2.font = Font(name=_XLSX_FONT, size=9, color="1E293B")
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        c2.border = _XLSX_BORDER
        ws6.row_dimensions[r].height = 30
        r += 1
    if not action_plan:
        ws6.merge_cells("A2:B2")
        c = ws6.cell(row=2, column=1, value="No action plan items for this run.")
        c.font = Font(name=_XLSX_FONT, italic=True, size=9, color="94A3B8")
    _xtitle(ws6, "AI-GENERATED ACTION PLAN", "C9A227")
    _xautofit(ws6, [10, 100])

    # ── SHEET 7 — Environment ─────────────────────────────────────────────
    ws7 = wb.create_sheet("Environment")
    ws7.sheet_view.showGridLines = False
    _xh(ws7, 1, ["Property", "Value"])
    ws7.freeze_panes = "A2"
    total_ms = 0
    for t in tests:
        d = str(t.get('duration', '0'))
        try:
            if d.endswith('ms'):
                total_ms += float(d.replace('ms', '') or 0)
            elif d.endswith('s'):
                total_ms += float(d.replace('s', '') or 0) * 1000
        except Exception:
            pass
    env_items = [
        ("Framework", framework),
        ("Authentication Method", auth_method),
        ("Base URL", url),
        ("Total Tests", str(total)),
        ("Execution Time", f'{total_ms/1000:.2f}s' if total_ms else 'N/A'),
        ("NexTest Version", generation_data.get('nextest_version', '1.0.0')),
        ("Generated", datetime.now().strftime('%Y-%m-%d %H:%M')),
    ]
    r = 2
    for lbl, val in env_items:
        c1 = ws7.cell(row=r, column=1, value=lbl)
        c1.font = Font(name=_XLSX_FONT, bold=True, size=9, color="EAB308")
        c1.fill = PatternFill("solid", fgColor="F8FAFC")
        c1.border = _XLSX_BORDER
        c2 = ws7.cell(row=r, column=2, value=val)
        c2.font = Font(name=_XLSX_FONT, size=9, color="1E293B")
        c2.border = _XLSX_BORDER
        r += 1
    _xtitle(ws7, "ENVIRONMENT & EXECUTION INFO", "EAB308")
    _xautofit(ws7, [24, 40])

    # ── SHEET 8 — Certificate ─────────────────────────────────────────────
    ws8 = wb.create_sheet("Certificate")
    ws8.sheet_view.showGridLines = False
    _xh(ws8, 1, ["Property", "Value"])
    ws8.freeze_panes = "A2"

    if critical_fail:
        verdict_text = (f'Internal Regression Test FAILED — critical checks did not pass on {url}. '
                         f'This application is NOT ready for the next deployment until resolved.')
        verdict_color = "EF4444"
        risk_final = "HIGH"
    elif fail_count > 0:
        verdict_text = (f'Internal Regression Test passed with {fail_count} non-critical issue(s) on {url}. '
                         f'Review before the next deployment.')
        verdict_color = "F59E0B"
        risk_final = "MEDIUM"
    else:
        verdict_text = f'Internal Regression Test PASSED — all {pass_count} test(s) succeeded on {url}.'
        verdict_color = "10B981"
        risk_final = "LOW"

    ws8.merge_cells("A1:B1")
    c = ws8.cell(row=1, column=1, value="Final Regression Verdict")
    c.font = Font(name=_XLSX_FONT, bold=True, size=13, color=verdict_color)
    r = 2
    n_rows_v = max(2, math.ceil(len(verdict_text) / 95) + 1)
    ws8.merge_cells(start_row=r, start_column=1, end_row=r + n_rows_v - 1, end_column=2)
    c = ws8.cell(row=r, column=1, value=verdict_text)
    c.font = Font(name=_XLSX_FONT, size=9, color=verdict_color)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.fill = PatternFill("solid", fgColor=("D1FAE5" if risk_final == "LOW" else "FFFBEB" if risk_final == "MEDIUM" else "FEE2E2"))
    r += n_rows_v
    ws8.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws8.cell(row=r, column=1,
                 value=f"Regression Score: {score}/100  |  Risk Level: {risk_final}  |  Deployment Readiness: {score_label}")
    c.font = Font(name=_XLSX_FONT, bold=True, size=10, color="1E293B")
    r += 2

    cert_details = [
        ("Validation Date", datetime.now().strftime('%Y-%m-%d  %H:%M')),
        ("Framework", framework),
        ("Authentication", auth_method),
        ("Base URL", url),
        ("Execution Time", f'{total_ms/1000:.2f}s' if total_ms else 'N/A'),
        ("Overall Grade", grade),
        ("AI Validation Status", "Verified by NexTest AI"),
    ]
    for lbl, val in cert_details:
        c1 = ws8.cell(row=r, column=1, value=lbl)
        c1.font = Font(name=_XLSX_FONT, bold=True, size=9, color="F97316")
        c1.fill = PatternFill("solid", fgColor="F8FAFC")
        c1.border = _XLSX_BORDER
        c2 = ws8.cell(row=r, column=2, value=val)
        c2.font = Font(name=_XLSX_FONT, size=9, color="1E293B")
        c2.border = _XLSX_BORDER
        r += 1

    r += 2
    ws8.merge_cells(f"A{r}:B{r}")
    c = ws8.cell(row=r, column=1, value=f"{url}  —  {score}/100  ({score_label})")
    c.font = Font(name=_XLSX_FONT, bold=True, size=12, color=score_color)
    c.alignment = Alignment(horizontal="center")
    r += 2
    ws8.merge_cells(f"A{r}:B{r}")
    c = ws8.cell(row=r, column=1, value=f"Validated by NexTest AI  •  {datetime.now().strftime('%Y-%m-%d')}")
    c.font = Font(name=_XLSX_FONT, italic=True, size=8, color="94A3B8")
    c.alignment = Alignment(horizontal="center")

    _xtitle(ws8, "CERTIFICATE OF INTERNAL REGRESSION VALIDATION", "EAB308")
    _xautofit(ws8, [26, 40])

    # ── Ordre final ──────────────────────────────────────────────────────
    order = ["Overview", "Test Scenarios", "Results by Category", "Detailed Results",
             "AI Recommendations", "Action Plan", "Environment", "Certificate"]
    wb._sheets = [wb[name] for name in order if name in wb.sheetnames]
    wb.active = 0

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()